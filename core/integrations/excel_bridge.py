"""Stable Excel interchange for takeoff/BOQ data."""
from __future__ import annotations
from pathlib import Path
def export_rows(rows,path):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title="BOQ"
    headers=["ردیف","کد","شرح","مقدار","واحد","بهای واحد","مبلغ"]
    ws.append(headers)
    for i,r in enumerate(rows,1): ws.append([i,r.get("price_code",""),r.get("description",""),r.get("quantity",0),r.get("unit",""),r.get("unit_price",0),r.get("total",0)])
    wb.save(path); return Path(path)
def import_rows(path):
    from openpyxl import load_workbook
    ws=load_workbook(path,data_only=True).active
    rows=[]
    for values in ws.iter_rows(min_row=2,values_only=True):
        if not any(v is not None for v in values): continue
        rows.append({"price_code":values[1] or "","description":values[2] or "","quantity":float(values[3] or 0),
                     "unit":values[4] or "","unit_price":float(values[5] or 0),"total":float(values[6] or 0)})
    return rows
