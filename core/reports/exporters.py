"""Professional offline report exporters with RTL-aware Excel/PDF output."""
from __future__ import annotations
from pathlib import Path
from typing import Any
import re
from xml.sax.saxutils import escape


def _summary_items(summary):
    summary = summary or {}
    cost = summary.get("cost", summary) if isinstance(summary, dict) else {}
    items = []
    labels = (("base", "مبلغ پایه"), ("grand_total", "مبلغ نهایی"), ("line_count", "تعداد ردیف"),
              ("active_line_count", "ردیف‌های فعال"))
    for key, label in labels:
        if key in cost:
            items.append((label, cost.get(key)))
    for name, value in (cost.get("factors", {}) or {}).items():
        items.append((str(name), value))
    for name, value in (summary.get("amount_by_group", {}) or {}).items():
        items.append((f"گروه: {name}", value))
    statement = summary.get("statement", {}) if isinstance(summary, dict) else {}
    statement_labels = (
        ("gross_current", "مبلغ ناخالص این دوره"), ("retention", "کسور تضمین"),
        ("advance_recovery", "استهلاک پیش‌پرداخت"), ("taxable_current", "مبلغ مشمول مالیات"),
        ("tax", "مالیات"), ("insurance", "بیمه"), ("payable_current", "قابل پرداخت این دوره"),
        ("previous_paid", "پرداختی قبلی"), ("balance_after_current", "مانده پس از این دوره"),
    )
    for key, label in statement_labels:
        if key in statement:
            items.append((label, statement.get(key)))
    if "period_no" in summary:
        items.insert(0, ("شماره صورت‌وضعیت", summary["period_no"]))
    return items


def _columns(rows, columns=None):
    rows = list(rows)
    if columns:
        return list(columns)
    if not rows:
        return [("description", "شرح")]
    keys = list(rows[0].keys())
    return [(key, key.replace("_", " ")) for key in keys]


def _cell(value):
    if value is None:
        return ""
    return str(value)


def export_csv(rows, path, title="StructuralPro", summary=None, metadata=None, columns=None):
    import csv
    rows = list(rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = _columns(rows, columns)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([title])
        if metadata:
            for key, value in metadata.items():
                w.writerow([key, value])
        w.writerow([])
        w.writerow([label for _, label in cols])
        w.writerows([_cell(row.get(key)) for key, _ in cols] for row in rows)
    return path


def export_xlsx(rows, path, title="StructuralPro", summary=None, metadata=None, columns=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    rows = list(rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = _columns(rows, columns)
    wb = Workbook()
    detail = wb.active
    detail.title = "ریز متره"
    detail.sheet_view.rightToLeft = True
    summary_ws = wb.create_sheet("خلاصه")
    summary_ws.sheet_view.rightToLeft = True
    info_ws = wb.create_sheet("اطلاعات")
    info_ws.sheet_view.rightToLeft = True

    detail.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(cols)))
    detail.cell(1, 1, title)
    detail.cell(1, 1).font = Font(bold=True, size=16)
    detail.cell(1, 1).alignment = Alignment(horizontal="right", readingOrder=2)
    detail.append([])
    detail.cell(3, 1, "خلاصه تجاری")
    detail.cell(3, 1).font = Font(bold=True)
    detail.cell(3, 1).alignment = Alignment(horizontal="right", readingOrder=2)
    summary_items = _summary_items(summary)
    for i, (label, value) in enumerate(summary_items, 4):
        detail.cell(i, 1, label)
        detail.cell(i, 2, value)
    header_row = max(6, 4 + len(summary_items))
    detail.cell(header_row, 1, "ریز متره و برآورد")
    detail.cell(header_row, 1).font = Font(bold=True)
    header_row += 1
    detail.append([label for _, label in cols])
    for row in rows:
        detail.append([row.get(key, "") for key, _ in cols])
    thin = Side(style="thin")
    for cell in detail[header_row]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right", readingOrder=2, vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in detail.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", readingOrder=2, vertical="top", wrap_text=True)
    detail.freeze_panes = f"A{header_row + 1}"
    detail.auto_filter.ref = f"A{header_row}:{get_column_letter(max(1, len(cols)))}{detail.max_row}"
    for idx, (key, label) in enumerate(cols, 1):
        width = max(12, min(42, max([len(str(label))] + [len(_cell(r.get(key))) for r in rows]) + 2))
        detail.column_dimensions[get_column_letter(idx)].width = width

    summary_ws["A1"] = title
    summary_ws["A1"].font = Font(bold=True, size=16)
    summary_ws["A3"], summary_ws["B3"] = "شرح", "مقدار"
    for i, (label, value) in enumerate(_summary_items(summary), 4):
        summary_ws.cell(i, 1, label)
        summary_ws.cell(i, 2, value)
    for cell in summary_ws[3]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right", readingOrder=2)
    for row in summary_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right", readingOrder=2, vertical="top", wrap_text=True)
    summary_ws.column_dimensions["A"].width = 32
    summary_ws.column_dimensions["B"].width = 24

    info_ws["A1"], info_ws["B1"] = "فیلد", "مقدار"
    info_ws["A1"].font = info_ws["B1"].font = Font(bold=True)
    info_ws["A2"], info_ws["B2"] = "نام گزارش", title
    for i, (key, value) in enumerate((metadata or {}).items(), 3):
        info_ws.cell(i, 1, key)
        info_ws.cell(i, 2, value)
    for row in info_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="right", readingOrder=2, wrap_text=True)
    info_ws.column_dimensions["A"].width = 28
    info_ws.column_dimensions["B"].width = 48
    wb.save(path)
    return path


def export_docx(rows, path, title="StructuralPro", summary=None, metadata=None, columns=None):
    from docx import Document
    from docx.enum.table import WD_TABLE_DIRECTION
    rows = list(rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = _columns(rows, columns)
    d = Document()
    d.add_heading(title, 0)
    d.add_heading("خلاصه تجاری", level=1)
    table = d.add_table(rows=1, cols=2)
    try:
        table.direction = WD_TABLE_DIRECTION.RTL
    except Exception:
        pass
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "شرح", "مقدار"
    for label, value in _summary_items(summary):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(label), _cell(value)
    d.add_heading("ریز متره و برآورد", level=1)
    t = d.add_table(rows=1, cols=len(cols))
    try:
        t.direction = WD_TABLE_DIRECTION.RTL
    except Exception:
        pass
    for i, (_, label) in enumerate(cols):
        t.rows[0].cells[i].text = str(label)
    for row in rows:
        cells = t.add_row().cells
        for i, (key, _) in enumerate(cols):
            cells[i].text = _cell(row.get(key))
    d.save(path)
    return path


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\tahoma.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                pdfmetrics.registerFont(TTFont("StructuralProUnicode", candidate))
                return "StructuralProUnicode"
            except Exception:
                pass
    raise RuntimeError("فونت Unicode فارسی برای خروجی PDF در دسترس نیست")


_PERSIAN_TEXT = re.compile(r"[\u0600-\u06ff]")
_PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def _pdf_text(value):
    """Return safe, visually ordered Persian text for ReportLab paragraphs."""
    text = _cell(value)
    if _PERSIAN_TEXT.search(text):
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
        except ImportError as exc:
            raise RuntimeError(
                "موتور شکل‌دهی و راست‌به‌چپ فارسی برای خروجی PDF نصب نیست"
            ) from exc
        text = get_display(arabic_reshaper.reshape(text), base_dir="R")
    return escape(text)


def _pdf_page_footer(canvas, doc, font):
    canvas.saveState()
    canvas.setFont(font, 8)
    label = f"صفحه {doc.page}".translate(_PERSIAN_DIGITS)
    canvas.drawRightString(doc.pagesize[0] - doc.rightMargin, 16, _pdf_text(label))
    canvas.restoreState()


def export_pdf(rows, path, title="StructuralPro", summary=None, metadata=None, columns=None):
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors
    rows = list(rows)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cols = _columns(rows, columns)
    font = _pdf_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("StructuralProBody", parent=styles["BodyText"], fontName=font,
                          alignment=TA_RIGHT, leading=13, fontSize=8)
    heading = ParagraphStyle("StructuralProHeading", parent=styles["Heading2"], fontName=font,
                             alignment=TA_RIGHT)
    title_style = ParagraphStyle("StructuralProTitle", parent=styles["Title"], fontName=font,
                                 alignment=TA_RIGHT)
    # Landscape keeps professional BOQ tables readable while remaining printable on A4.
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4),
                            rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=34)
    story = [Paragraph(_pdf_text(title), title_style), Spacer(1, 10),
             Paragraph(_pdf_text("خلاصه تجاری"), heading)]
    summary_rows = _summary_items(summary)
    if summary_rows:
        data = [[Paragraph(_pdf_text("شرح"), body), Paragraph(_pdf_text("مقدار"), body)]]
        data += [[Paragraph(_pdf_text(a), body), Paragraph(_pdf_text(b), body)] for a, b in summary_rows]
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([("GRID", (0,0), (-1,-1), .4, colors.grey),
                                   ("ALIGN", (0,0), (-1,-1), "RIGHT"),
                                   ("VALIGN", (0,0), (-1,-1), "TOP")]))
        story += [table, Spacer(1, 12)]
    story.append(Paragraph(_pdf_text("ریز متره و برآورد"), heading))
    # Reverse display order for RTL so the first logical field is on the right.
    display_cols = list(reversed(cols))
    data = [[Paragraph(_pdf_text(label), body) for _, label in display_cols]]
    for row in rows:
        data.append([Paragraph(_pdf_text(row.get(key)), body) for key, _ in display_cols])
    if not rows:
        data.append([Paragraph(_pdf_text("ردیفی برای نمایش وجود ندارد"), body)])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([("GRID", (0,0), (-1,-1), .35, colors.grey),
                               ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
                               ("ALIGN", (0,0), (-1,-1), "RIGHT"),
                               ("VALIGN", (0,0), (-1,-1), "TOP")]))
    story.append(table)
    footer = lambda canvas, current_doc: _pdf_page_footer(canvas, current_doc, font)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return path

