"""P97 — production-grade Excel/CSV/JSON pricebook row extraction.

Evidence-first: header ambiguity fails closed; every extracted row carries file hash
and sheet provenance. No inferred price is created.
"""
from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from hashlib import sha256
import csv, json, re

@dataclass(frozen=True)
class NormalizedPriceRow:
    year:int
    discipline:str
    item_code:str
    description:str
    unit:str
    unit_price:float|None
    source_sha256:str
    source_file:str
    source_sheet:str=""

HEADER_ALIASES={
    "item_code":("item_code","code","item code","ردیف","کد","کد ردیف","شماره ردیف"),
    "description":("description","شرح","شرح کار","شرح عملیات"),
    "unit":("unit","واحد"),
    "unit_price":("unit_price","price","rate","بهای واحد","بهای واحد (ریال)","بهای واحد(ریال)","مبلغ واحد"),
    "chapter":("chapter","فصل","شماره فصل","عنوان فصل"),
}

def _text(v): return "" if v is None else str(v).strip()

def _norm_header(v):
    s=_text(v).lower()
    s=s.replace("ي","ی").replace("ك","ک").replace("\u200c","")
    s=re.sub(r"[\s_\-–—]+","",s)
    s=s.replace("(","").replace(")","")
    return s

_ALIAS_INDEX={_norm_header(a):k for k,vals in HEADER_ALIASES.items() for a in vals}

def _price(v):
    s=_text(v).replace(",","").replace("٬","").replace("،","").replace(" ","")
    if not s: return None
    trans=str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩","01234567890123456789")
    try: return float(s.translate(trans))
    except ValueError: return None

def detect_header_row(values, scan_limit=30):
    candidates=[]
    for idx,row in enumerate(values[:scan_limit]):
        mapped=[_ALIAS_INDEX.get(_norm_header(v)) for v in row]
        present={x for x in mapped if x}
        score=sum(x in present for x in ("item_code","description","unit","unit_price"))
        if score>=3:
            candidates.append((score,idx,mapped))
    if not candidates:
        raise ValueError("pricebook header row not detected")
    best_score=max(x[0] for x in candidates)
    best=[x for x in candidates if x[0]==best_score]
    if len(best)!=1:
        raise ValueError("ambiguous pricebook header row")
    return best[0][1],best[0][2]

def _row(raw,mapping,year,discipline,digest,name,sheet=""):
    def get(key):
        try: return raw[mapping.index(key)]
        except ValueError: return None
    code=_text(get("item_code")); desc=_text(get("description")); unit=_text(get("unit"))
    price=_price(get("unit_price"))
    if not code and not desc: return None
    return NormalizedPriceRow(year,discipline,code,desc,unit,price,digest,name,sheet)

def extract_rows(values,year,discipline,digest,name,sheet=""):
    if not values: return []
    header_idx,mapping=detect_header_row(values)
    out=[]
    for raw in values[header_idx+1:]:
        x=_row(list(raw),mapping,year,discipline,digest,name,sheet)
        if x: out.append(x)
    return out

def extract_csv(path,year,discipline):
    p=Path(path); digest=sha256(p.read_bytes()).hexdigest()
    with p.open("r",encoding="utf-8-sig",newline="") as fh:
        return extract_rows(list(csv.reader(fh)),year,discipline,digest,p.name,"")

def extract_json(path,year,discipline):
    p=Path(path); digest=sha256(p.read_bytes()).hexdigest()
    data=json.loads(p.read_text(encoding="utf-8"))
    if isinstance(data,dict): data=data.get("rows",[])
    if not isinstance(data,list): raise ValueError("JSON pricebook rows must be a list")
    if not data: raise ValueError("no JSON pricebook rows")
    keys=list(data[0].keys())
    return [x for r in data for x in [_row([r.get(k) for k in keys],
        [_ALIAS_INDEX.get(_norm_header(k)) for k in keys],year,discipline,digest,p.name,"")] if x]

def extract_xlsx(path,year,discipline):
    from openpyxl import load_workbook
    p=Path(path); digest=sha256(p.read_bytes()).hexdigest(); out=[]
    wb=load_workbook(p,data_only=True,read_only=True)
    for ws in wb.worksheets:
        values=list(ws.iter_rows(values_only=True))
        if not values: continue
        if ws.sheet_state=="hidden" and not any(any(v is not None for v in r) for r in values):
            continue
        # Real pricebooks commonly contain cover/metadata sheets before the
        # tabular sheets. Skip sheets with no recognizable header, but keep
        # ambiguity fail-closed so malformed tables are never guessed.
        try:
            detect_header_row(values)
        except ValueError as exc:
            if str(exc) == "pricebook header row not detected":
                continue
            raise
        out.extend(extract_rows(values,year,discipline,digest,p.name,ws.title))
    if not out: raise ValueError("no parseable XLSX pricebook rows")
    return out

def validate_rows(rows):
    if not rows: raise ValueError("no pricebook rows extracted")
    bad=[r.source_file for r in rows if not r.item_code and not r.description]
    if bad: raise ValueError("row without code and description")
    return {"rows":len(rows),"priced_rows":sum(r.unit_price is not None for r in rows),
            "unpriced_rows":sum(r.unit_price is None for r in rows),
            "source_hashes":sorted({r.source_sha256 for r in rows}),
            "sheets":sorted({r.source_sheet for r in rows if r.source_sheet})}
