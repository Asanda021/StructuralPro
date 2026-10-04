"""P68 — import real pricebook files and connect rates to takeoff rows."""
from __future__ import annotations
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
import csv, json, zipfile

@dataclass(frozen=True)
class PriceRow:
    year:int; chapter:str; item_code:str; description:str; unit:str; rate:float; currency:str="IRR"

@dataclass(frozen=True)
class PricedTakeoff:
    item_code:str; quantity:float; unit:str; rate:float; currency:str; amount:float; pricebook_year:int

def _num(v):
    if isinstance(v,(int,float)): return float(v)
    s=str(v or "").strip().replace("٬","").replace(",","").replace("،","")
    trans=str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩","01234567890123456789")
    return float(s.translate(trans))

def parse_rows(records, year, currency="IRR"):
    out=[]
    for r in records:
        code=str(r.get("item_code") or r.get("کد") or r.get("ردیف") or "").strip()
        desc=str(r.get("description") or r.get("شرح") or "").strip()
        unit=str(r.get("unit") or r.get("واحد") or "").strip()
        raw=r.get("rate",r.get("بهای واحد",r.get("بهای واحد (ریال)",None)))
        if not code and not desc and raw in (None,""): continue
        if not (code and desc and unit and raw not in (None,"")): raise ValueError("incomplete pricebook row")
        rate=_num(raw)
        if rate < 0: raise ValueError("negative rate")
        out.append(PriceRow(year,str(r.get("chapter") or r.get("فصل") or "").strip(),code,desc,unit,rate,currency))
    if not out: raise ValueError("no pricebook rows")
    return tuple(out)

def load_json(path,year,currency="IRR"):
    return parse_rows(json.loads(Path(path).read_text(encoding="utf-8")),year,currency)

def load_csv(path,year,currency="IRR"):
    with open(path,encoding="utf-8-sig",newline="") as f: return parse_rows(csv.DictReader(f),year,currency)

def load_xls(path,year,currency="IRR"):
    try:
        import pandas as pd
    except ImportError as e:
        raise RuntimeError("pandas is required for XLS ingestion") from e
    sheets=pd.read_excel(path, sheet_name=None, dtype=object)
    rows=[]
    for frame in sheets.values(): rows.extend(frame.to_dict("records"))
    return parse_rows(rows, year, currency)

def load_archive(path, year, currency="IRR"):
    p=Path(path)
    if p.suffix.lower()==".zip":
        opener=lambda: zipfile.ZipFile(p)
    else:
        try: import rarfile
        except ImportError as e: raise RuntimeError("RAR ingestion requires rarfile") from e
        opener=lambda: rarfile.RarFile(p)
    with opener() as archive:
        names=archive.namelist()
        supported=[x for x in names if Path(x).suffix.lower() in {".xlsx",".xls",".csv",".json"}]
        if not supported: raise ValueError("archive contains no supported pricebook file")
        name=supported[0]
        tmp=p.parent/(p.stem+"_"+Path(name).name)
        tmp.write_bytes(archive.read(name))
        try: return load_supported(tmp,year,currency)
        finally: tmp.unlink(missing_ok=True)

def load_supported(path,year,currency="IRR"):
    ext=Path(path).suffix.lower()
    if ext==".xlsx": return load_xlsx(path,year,currency)
    if ext==".xls": return load_xls(path,year,currency)
    if ext==".csv": return load_csv(path,year,currency)
    if ext==".json": return load_json(path,year,currency)
    if ext in {".zip",".rar"}: return load_archive(path,year,currency)
    raise ValueError(f"unsupported pricebook extension: {ext}")
def load_xlsx(path,year,currency="IRR"):
    try:
        from openpyxl import load_workbook
    except ImportError as e: raise RuntimeError("openpyxl is required for XLSX ingestion") from e
    wb=load_workbook(path,data_only=True,read_only=True)
    rows=[]
    for ws in wb.worksheets:
        headers=[str(c.value or "").strip() for c in next(ws.iter_rows())]
        for cells in ws.iter_rows(min_row=2,values_only=True):
            row=dict(zip(headers,cells))
            rows.append(row)
    return parse_rows(rows,year,currency)

def index(rows):
    result={}
    for r in rows:
        k=(r.item_code,r.unit)
        if k in result and result[k].rate != r.rate: raise ValueError("conflicting pricebook rates")
        result[k]=r
    return result

def price_takeoff(takeoff_rows, price_rows):
    rates=index(price_rows); out=[]
    for row in takeoff_rows:
        code=str(row["item_code"]).strip(); unit=str(row["unit"]).strip(); qty=_num(row["quantity"])
        if qty < 0: raise ValueError("negative takeoff quantity")
        rate=rates.get((code,unit))
        if rate is None: raise KeyError(f"missing pricebook rate: {code}/{unit}")
        out.append(PricedTakeoff(code,qty,unit,rate.rate,rate.currency,qty*rate.rate,rate.year))
    return tuple(out)

def fingerprint(rows):
    payload="|".join(sorted(f"{r.year}|{r.chapter}|{r.item_code}|{r.description}|{r.unit}|{r.rate}|{r.currency}" for r in rows))
    return sha256(payload.encode()).hexdigest()


def coverage(rows, expected_year=None):
    if not rows: raise ValueError("no rows for coverage")
    years=sorted({r.year for r in rows})
    codes=sorted({r.item_code for r in rows})
    return {"row_count":len(rows),"years":years,"item_code_count":len(codes),"fingerprint":fingerprint(rows),"year_ok":expected_year is None or years==[expected_year]}
