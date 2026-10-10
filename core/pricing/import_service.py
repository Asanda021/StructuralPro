"""Fail-closed price-list import with provenance and user-import support.

Supports CSV and real Excel workbooks (XLSX/XLSM). Excel data is normalized to
StructuralPro's canonical price-item schema before it reaches the catalog.
"""
from __future__ import annotations

from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any
import csv
from io import StringIO

from core.pricing.catalog import PriceCatalog, PriceItem
from core.pricing.source_registry import PriceSource, PriceSourceRegistry


@dataclass(frozen=True)
class ImportReceipt:
    source_id: str
    year: int
    discipline: str
    filename: str
    sha256: str
    rows: int
    verified_source: bool
    imported_at: str
    format: str = "csv"

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


class PricebookImportService:
    """Import official/user datasets without inventing source provenance."""

    _ALIASES = {
        "year": {"year", "سال", "سال فهرست بها", "سال فهرست‌بها"},
        "group": {"group", "رشته", "دسته", "گروه", "دیسپلین", "رشته کاری"},
        "chapter": {"chapter", "فصل", "فصل کاری", "فصل فهرست بها", "فصل فهرست‌بها"},
        "code": {"code", "کد", "شماره ردیف", "شماره ردیف فهرست بها", "شماره فهرست بها", "شماره فهرست‌بها", "ردیف"},
        "description": {"description", "شرح", "شرح ردیف", "شرح عملیات"},
        "unit": {"unit", "واحد", "واحد اندازه گیری", "واحد اندازه‌گیری"},
        "unit_price": {"unit_price", "unit price", "بهای واحد", "بهای واحد (ریال)", "قیمت واحد", "قیمت", "بها"},
        "analysis": {"analysis", "تجزیه", "تجزیه بها", "آنالیز"},
        "notes": {"notes", "یادداشت", "توضیحات", "توضیح"},
    }

    def __init__(self, catalog: PriceCatalog | None = None, registry: PriceSourceRegistry | None = None):
        self.catalog = catalog or PriceCatalog()
        self.registry = registry or PriceSourceRegistry()

    @staticmethod
    def _norm(value: Any) -> str:
        return " ".join(str(value or "").strip().replace("‌", " ").split()).casefold()

    @classmethod
    def _map_headers(cls, headers: list[Any]) -> dict[str, int]:
        normalized = [cls._norm(x) for x in headers]
        mapping: dict[str, int] = {}
        for field, aliases in cls._ALIASES.items():
            aliases_norm = {cls._norm(x) for x in aliases}
            for idx, header in enumerate(normalized):
                if header in aliases_norm:
                    mapping[field] = idx
                    break
        return mapping

    @staticmethod
    def _number(value: Any) -> float:
        raw = str(value or "").strip()
        trans = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
        raw = raw.translate(trans).replace(",", "").replace("٬", "").replace(" ", "")
        if not raw:
            return 0.0
        return float(raw)

    @classmethod
    def _unit_price(cls, value: Any) -> float:
        # Blank price is missing data, not an explicitly recorded zero price.
        # Zero is permitted but must be present in the source cell.
        if value is None or isinstance(value, bool) or (
            isinstance(value, str) and not value.strip()
        ):
            raise ValueError("بهای واحد خالی یا نامعتبر است؛ قیمت صفر باید صریح ثبت شود.")
        return cls._number(value)

    @classmethod
    def _rows_from_excel(cls, path: Path, *, fallback_year: int) -> list[PriceItem]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("برای ورود Excel باید openpyxl نصب باشد.") from exc
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                rows = ws.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    continue
                mapping = cls._map_headers(list(headers))
                required = {"code", "description", "unit", "unit_price"}
                if not required.issubset(mapping):
                    continue
                out: list[PriceItem] = []
                for raw in rows:
                    values = list(raw)
                    def val(field: str, default: Any = ""):
                        idx = mapping.get(field)
                        return values[idx] if idx is not None and idx < len(values) else default
                    code = str(val("code") or "").strip()
                    desc = str(val("description") or "").strip()
                    unit = str(val("unit") or "").strip()
                    price_value = val("unit_price")
                    if not any(str(value or "").strip() for value in (code, desc, unit, price_value)):
                        continue  # genuinely empty spreadsheet row
                    if not code or not desc or not unit:
                        raise ValueError("ردیف ناقص Excel: کد، شرح و واحد باید مشخص باشند.")
                    year_value = val("year", fallback_year)
                    group = str(val("group", "")).strip()
                    chapter = str(val("chapter", "")).strip()
                    price = cls._unit_price(price_value)
                    out.append(PriceItem(
                        year=int(cls._number(year_value)) if year_value else 0,
                        group=group, chapter=chapter, code=code,
                        description=desc, unit=unit, unit_price=price,
                        analysis=str(val("analysis", "") or ""),
                        notes=str(val("notes", "") or ""),
                    ))
                if out:
                    return out
            raise ValueError(
                "ساختار فایل Excel قابل تشخیص نیست. ستون‌های ضروری: کد/شماره ردیف، شرح، واحد و بهای واحد."
            )
        finally:
            wb.close()

    @classmethod
    def _rows_from_pdf(cls, path: Path, *, fallback_year: int) -> list[PriceItem]:
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("برای ورود PDF باید PyMuPDF نصب باشد.") from exc
        text_parts=[]
        with fitz.open(path) as doc:
            for page in doc:
                text_parts.append(page.get_text("text") or "")
        text="\n".join(text_parts)
        if not text.strip():
            raise ValueError("PDF اسکن‌شده یا تصویری است و برای جلوگیری از حدس‌زدن رد شد.")
        out=[]
        for line in text.splitlines():
            cols=[x.strip() for x in line.split("|") if x.strip()]
            if len(cols)<4: continue
            code,desc,unit,raw_price=cols[0],cols[1],cols[2],cols[3]
            try: price=cls._unit_price(raw_price)
            except ValueError: continue
            if code and desc and unit and price>=0:
                out.append(PriceItem(year=fallback_year,group="",chapter="",code=code,description=desc,unit=unit,unit_price=price))
        if not out:
            raise ValueError("PDF قابل تشخیص نیست؛ PDF باید متن واقعی با ستون‌های کد، شرح، واحد و بهای واحد داشته باشد.")
        return out

    @classmethod
    def _rows_from_csv(cls, path: Path, *, fallback_year: int) -> list[PriceItem]:
        text = path.read_text(encoding="utf-8-sig")
        reader = csv.reader(StringIO(text))
        headers = next(reader, None)
        if not headers:
            return []
        mapping = cls._map_headers(headers)
        required = {"code", "description", "unit", "unit_price"}
        if not required.issubset(mapping):
            raise ValueError("ساختار CSV قابل تشخیص نیست. ستون‌های ضروری: کد، شرح، واحد و بهای واحد.")
        out: list[PriceItem] = []
        for raw in reader:
            values = list(raw)
            def val(field: str, default: Any = ""):
                idx = mapping.get(field)
                return values[idx] if idx is not None and idx < len(values) else default
            out.append(PriceItem(
                year=int(cls._number(val("year", fallback_year)) or fallback_year),
                group=str(val("group", "") or ""),
                chapter=str(val("chapter", "") or ""),
                code=str(val("code", "") or "").strip(),
                description=str(val("description", "") or "").strip(),
                unit=str(val("unit", "") or "").strip(),
                unit_price=cls._unit_price(val("unit_price")),
                analysis=str(val("analysis", "") or ""),
                notes=str(val("notes", "") or ""),
            ))
        return out

    def _load_items(self, path: Path, *, fallback_year: int) -> tuple[list[PriceItem], str]:
        suffix = path.suffix.casefold()
        if suffix in {".xlsx", ".xlsm"}:
            return self._rows_from_excel(path, fallback_year=fallback_year), "excel"
        if suffix == ".csv":
            return self._rows_from_csv(path, fallback_year=fallback_year), "csv"
        if suffix == ".pdf":
            return self._rows_from_pdf(path, fallback_year=fallback_year), "pdf"
        raise ValueError("فرمت پشتیبانی‌شده برای فهرست‌بها: XLSX، XLSM، CSV یا PDF متنی")

    def inspect(self, path: str | Path, *, year: int, discipline: str = "building",
                source_id: str = "user-import") -> dict[str, Any]:
        p = Path(path)
        if not p.exists():
            raise FileNotFoundError(p)
        raw_bytes = p.read_bytes()
        items, file_format = self._load_items(p, fallback_year=year)
        errors: list[str] = []
        seen_keys: set[tuple[int, str]] = set()
        for item in items:
            try:
                validated = self.catalog._validate_item(item)
                key = (validated.year, validated.code)
                if key in seen_keys:
                    errors.append(f"{validated.code}: کد تکراری برای سال {validated.year}")
                seen_keys.add(key)
            except (ValueError, TypeError) as exc:
                errors.append(f"{item.code or '?'}: {exc}")
        if not items:
            errors.append("هیچ ردیف قابل استفاده‌ای در فایل پیدا نشد.")
        source = self.registry.get(year, discipline, source_id)
        return {
            "path": str(p), "filename": p.name, "year": year,
            "discipline": discipline, "source_id": source_id,
            "sha256": sha256(raw_bytes).hexdigest(), "rows": len(items),
            "valid": not errors, "errors": errors[:20],
            "verified_source": bool(source and source.verified and
                                    self.registry.verify_record(source, sha256(raw_bytes).hexdigest())),
            "format": file_format,
        }

    def import_file(self, path: str | Path, *, year: int, discipline: str = "building",
                    source_id: str = "user-import", require_verified_source: bool = False,
                    replace_year: bool = True) -> ImportReceipt:
        p = Path(path)
        info = self.inspect(p, year=year, discipline=discipline, source_id=source_id)
        if not info["valid"]:
            raise ValueError("اعتبارسنجی فهرست‌بها ناموفق بود: " + "; ".join(info["errors"][:5]))
        if require_verified_source and not info["verified_source"]:
            raise PermissionError(
                "ورود رسمی/تأییدشده نیازمند منبع ثبت‌شده و SHA-256 منطبق است؛ حالت ورود فایل کاربر آزاد است."
            )
        items, file_format = self._load_items(p, fallback_year=year)
        normalized = [
            PriceItem(
                year=item.year or int(year), group=item.group, chapter=item.chapter,
                code=item.code, description=item.description, unit=item.unit,
                unit_price=item.unit_price, analysis=item.analysis, notes=item.notes,
            ) for item in items
        ]
        # Validate the complete incoming batch before mutating the catalog.
        # A rejected row must not partially replace existing pricebook data.
        validated = [self.catalog._validate_item(item) for item in normalized]
        incoming_keys = [(item.year, item.code) for item in validated]
        if len(set(incoming_keys)) != len(incoming_keys):
            raise ValueError("کد تکراری در فهرست‌بهای ورودی برای یک سال وجود دارد.")
        existing = dict(self.catalog._items)
        if replace_year:
            # A single import must not erase unrelated disciplines in the same
            # year. Legacy catalog keys are only (year, code).
            scopes = {(item.year, item.group) for item in validated}
            existing = {
                key: item for key, item in existing.items()
                if (item.year, item.group) not in scopes
            }
        for item in validated:
            key = (item.year, item.code)
            prior = existing.get(key)
            if prior is not None and prior.group != item.group:
                raise ValueError(
                    "کد یکسان در رشته‌های متفاوت وجود دارد؛ ورود برای جلوگیری از جایگزینی ناخواسته متوقف شد."
                )
            existing[key] = item
        # Never silently discard or detach a user's custom prices. Require
        # explicit resolution before replacing a row with an override.
        changed_keys = {key for key in set(self.catalog._items) | set(existing)
                        if self.catalog._items.get(key) != existing.get(key)}
        protected = changed_keys.intersection(self.catalog._overrides)
        if protected:
            raise ValueError(
                "ردیف دارای قیمت سفارشی است؛ پیش از جایگزینی، تعارض قیمت سفارشی را تعیین تکلیف کنید."
            )
        # Preserve earlier audit records and append changes only after the
        # entire import has passed validation and conflict checks.
        previous = self.catalog._items
        self.catalog._items = existing
        for item in validated:
            key = (item.year, item.code)
            old = previous.get(key)
            if old is not None and old.unit_price != item.unit_price:
                self.catalog._history.setdefault(key, []).append({
                    "year": item.year, "code": item.code,
                    "old_price": old.unit_price, "new_price": item.unit_price,
                })
        return ImportReceipt(
            source_id=source_id, year=year, discipline=discipline, filename=p.name,
            sha256=info["sha256"], rows=len(normalized),
            verified_source=info["verified_source"],
            imported_at=datetime.now(timezone.utc).isoformat(), format=file_format,
        )
