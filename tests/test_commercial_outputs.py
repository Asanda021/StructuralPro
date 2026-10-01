"""Acceptance tests for commercial report generation and output formats."""
from core.platform.application import StructuralProApp
from openpyxl import load_workbook


def _seed(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("پروژه خروجی", "R1")
    app.add_takeoff(
        "R1", "building", "wall",
        length=10, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )
    app.recalculate_estimate("R1")
    return app


def test_report_export_formats(tmp_path):
    app = _seed(tmp_path)
    project = app.open_project("R1")
    rows = [
        {"کد": q.get("price_code"), "شرح": q.get("title"), "مقدار": q.get("amount"),
         "واحد": q.get("unit"), "قیمت واحد": q.get("unit_price")}
        for t in project["takeoffs"] for q in t["quantities"]
    ]
    report = __import__("core.reports.project_report", fromlist=["build_report"]).build_report(
        project["name"], rows, {"grand_total": project["estimate"]["cost"]["grand_total"]}
    )
    for fmt in ("xlsx", "pdf", "docx", "csv"):
        path = tmp_path / f"report.{fmt}"
        report.export(path, fmt)
        assert path.exists() and path.stat().st_size > 0


def test_report_service_uses_project_data(tmp_path):
    app = _seed(tmp_path)
    path = tmp_path / "report.xlsx"
    result = app.report("R1", "xlsx", path)
    assert result == path
    assert path.exists() and path.stat().st_size > 0


def test_xlsx_contains_commercial_summary(tmp_path):
    app = _seed(tmp_path)
    path = tmp_path / "summary.xlsx"
    app.report("R1", "xlsx", path)
    ws = load_workbook(path, data_only=True).active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "خلاصه تجاری" in values
    assert "مبلغ پایه" in values
    assert "مبلغ نهایی" in values
    assert "کد" in values
    assert "W001" in values

def test_payment_statement_calculates_period_and_deductions(tmp_path):
    app = _seed(tmp_path)
    project = app.open_project("R1")
    project["boq"][0]["previous_quantity"] = 2
    project["boq"][0]["current_quantity"] = 3
    app.store.save("R1", project)
    statement = app.build_statement(
        "R1", previous_paid=500, retention_rate=0.05,
        advance_recovery_rate=0.10, tax_rate=0.10, insurance_rate=0.02,
    )
    assert statement["lines"][0]["cumulative_quantity"] == 5
    assert statement["gross_current"] == 3600
    assert statement["retention"] == 180
    assert statement["advance_recovery"] == 360
    assert statement["taxable_current"] == 3060
    assert statement["tax"] == 306
    assert statement["insurance"] == 61.2
    assert statement["payable_current"] == 3304.8
    assert statement["balance_after_current"] == 2804.8


def test_statement_export_contains_period_columns(tmp_path):
    app = _seed(tmp_path)
    project = app.open_project("R1")
    project["boq"][0]["current_quantity"] = 4
    app.store.save("R1", project)
    path = tmp_path / "statement.xlsx"
    app.report("R1", "xlsx", path)
    ws = load_workbook(path, data_only=True).active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "قبلی" in values
    assert "این دوره" in values
    assert "تجمعی" in values
    assert "مبلغ این دوره" in values

def test_statement_periods_roll_previous_quantities_forward(tmp_path):
    app = _seed(tmp_path)
    p = app.open_project("R1")
    p["boq"][0]["current_quantity"] = 2
    app.store.save("R1", p)
    first = app.save_statement_period("R1", retention_rate=0.05)
    assert first["number"] == 1
    assert first["lines"][0]["previous_quantity"] == 0
    assert first["lines"][0]["current_quantity"] == 2
    assert first["lines"][0]["cumulative_quantity"] == 2

    p = app.open_project("R1")
    p["boq"][0]["current_quantity"] = 3
    app.store.save("R1", p)
    second = app.save_statement_period("R1", retention_rate=0.05)
    assert second["number"] == 2
    assert second["lines"][0]["previous_quantity"] == 2
    assert second["lines"][0]["current_quantity"] == 3
    assert second["lines"][0]["cumulative_quantity"] == 5
    assert len(app.statement_periods("R1")) == 2


def test_statement_period_accepts_explicit_quantities_and_rejects_duplicate(tmp_path):
    app = _seed(tmp_path)
    first = app.save_statement_period("R1", current_quantities={"W001": 1.5})
    assert first["lines"][0]["current_quantity"] == 1.5
    try:
        app.save_statement_period("R1", period_no=1)
        assert False, "duplicate period should fail"
    except ValueError as exc:
        assert "unique" in str(exc)

def test_selected_statement_period_is_exported_with_summary(tmp_path):
    app = _seed(tmp_path)
    first = app.save_statement_period("R1", current_quantities={"W001": 2}, retention_rate=0.05)
    app.save_statement_period("R1", current_quantities={"W001": 3}, retention_rate=0.05)
    path = tmp_path / "period-1.xlsx"
    app.report("R1", "xlsx", path, period_no=first["number"])
    ws = load_workbook(path, data_only=True).active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "شماره صورت‌وضعیت" in values
    assert 1 in values
    assert "مبلغ ناخالص این دوره" in values
    assert "کسور تضمین" in values
    assert "W001" in values


def test_report_rejects_missing_statement_period(tmp_path):
    app = _seed(tmp_path)
    path = tmp_path / "missing.xlsx"
    try:
        app.report("R1", "xlsx", path, period_no=9)
        assert False, "missing period should fail"
    except KeyError as exc:
        assert "statement period 9" in str(exc)
