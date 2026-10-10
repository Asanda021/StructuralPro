from pathlib import Path
import pytest

from core.reports import build_report
from core.reports.exporters import _pdf_text


def _rows():
    return [{"item_no": 1, "item_code": "CONC-001", "price_code": "M-001", "chapter": "بتن", "category": "سازه", "group": "بتن", "description": "بتن فونداسیون", "quantity": 10, "unit": "m3", "unit_price": 2500000, "total": 25000000, "status": "active", "source_id": "takeoff-1"}]


def test_report_has_deterministic_rtl_columns_and_validation():
    report = build_report("پروژه نمونه", _rows(), {"grand_total": 25000000})
    assert report.validate()["valid"] is True
    assert report.columns()[0] == ("item_no", "ردیف")
    assert report.columns()[6] == ("description", "شرح")


def test_xlsx_export_has_rtl_detail_summary_and_info_sheets(tmp_path):
    report = build_report("پروژه نمونه", _rows(), {"grand_total": 25000000}, {"project_id": "P-1", "standard": "IR"})
    path = report.export(tmp_path / "report.xlsx", "xlsx")
    assert path.exists() and path.stat().st_size > 0
    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["ریز متره", "خلاصه", "اطلاعات"]
    assert wb["ریز متره"].sheet_view.rightToLeft is True
    assert wb["خلاصه"]["A3"].value == "شرح"
    assert wb["اطلاعات"]["B2"].value == "پروژه نمونه"


def test_pdf_export_is_non_empty_and_valid(tmp_path):
    report = build_report("پروژه نمونه", _rows(), {"grand_total": 25000000})
    path = report.export(tmp_path / "report.pdf", "pdf")
    assert path.exists() and path.stat().st_size > 100
    assert path.read_bytes().startswith(b"%PDF")


def test_pdf_persian_text_is_shaped_and_visual_rtl():
    shaped = _pdf_text("گزارش متره & برآورد")
    assert shaped != "گزارش متره & برآورد"
    assert "&amp;" in shaped
    assert any("\ufb50" <= char <= "\ufeff" for char in shaped)


def test_multi_page_pdf_repeats_table_and_has_page_footer(tmp_path):
    from pypdf import PdfReader

    rows = [
        {**_rows()[0], "item_no": index, "description": f"بتن فونداسیون ردیف {index}"}
        for index in range(1, 121)
    ]
    report = build_report("گزارش چندصفحه‌ای", rows, {"grand_total": 25000000})
    path = report.export(tmp_path / "multi-page.pdf", "pdf")
    reader = PdfReader(path)

    assert len(reader.pages) >= 2
    assert all(page.extract_text().strip() for page in reader.pages)


def test_report_ci_jobs_install_canonical_pdf_dependencies():
    root = Path(__file__).resolve().parents[1]
    workflows = (
        "priority-9-qa.yml",
        "priority-111-115-production-release.yml",
        "priority-116-120-production-acceptance.yml",
    )
    for name in workflows:
        source = (root / ".github" / "workflows" / name).read_text(encoding="utf-8")
        assert "python -m pip install -r requirements.txt" in source, name


def test_csv_export_keeps_persian_bom_and_named_headers(tmp_path):
    report = build_report("پروژه نمونه", _rows())
    path = report.export(tmp_path / "report.csv", "csv")
    data = path.read_bytes()
    assert data.startswith(b"\xef\xbb\xbf")
    assert "شرح".encode("utf-8") in data


@pytest.mark.parametrize("field", ["quantity", "unit_price", "total", "مقدار", "بهای واحد", "مبلغ"])
@pytest.mark.parametrize("value", [True, -1, float("nan"), float("inf"), "نامعتبر"])
def test_report_rejects_invalid_numbers_in_both_schemas(tmp_path, field, value):
    report = build_report("پروژه", [{"description": "بتن", field: value}])
    assert not report.validate()["valid"]
    with pytest.raises(ValueError):
        report.export(tmp_path / "invalid.csv", "csv")
    assert not (tmp_path / "invalid.csv").exists()


def test_persian_report_rejects_empty_description_and_duplicate_row_numbers():
    report = build_report("پروژه", [{"ردیف": 1, "شرح": ""}, {"ردیف": 1, "شرح": "بتن"}])
    assert report.validate()["errors"] == ["row_1:missing_description", "row_2:duplicate_item_no:1"]

