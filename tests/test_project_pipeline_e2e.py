""""End-to-end acceptance for the desktop project's commercial pipeline."""
import pytest
from core.platform.application import StructuralProApp
from core.takeoff.estimate import build_estimate


# CI trigger: keep this acceptance test in the main verification path.
def test_project_pipeline_persists_estimate_and_snapshot(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("پروژه آزمایشی", "P1")
    app.add_takeoff(
        "P1", "building", "wall",
        length=10, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )
    app.add_takeoff(
        "P1", "building", "wall",
        length=5, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )

    project_before_estimate = app.open_project("P1")
    estimate = build_estimate(project_before_estimate["boq"], factors={"سربار": 0.10}, aggregate=False)
    assert estimate["summary"]["line_count"] == 1
    assert estimate["cost"]["base"] > 0
    assert estimate["cost"]["grand_total"] == pytest.approx(estimate["cost"]["base"] * 1.10)

    persisted_estimate = app.recalculate_estimate("P1")
    assert persisted_estimate["cost"]["grand_total"] == persisted_estimate["cost"]["base"]

    snapshot = app.build_commercial_snapshot("P1")
    assert snapshot["project_id"] == "P1"
    assert snapshot["estimate"]["summary"]["grand_total"] > 0
    assert len(snapshot["progress"]["lines"]) == 1

    saved = app.open_project("P1")
    assert saved["estimate"]["cost"]["grand_total"] == snapshot["estimate"]["cost"]["grand_total"]
    assert saved["boq"][0]["price_code"] == "W001"

def test_financial_dashboard_summarizes_contract_progress_and_statements(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("پروژه مالی", "D1")
    app.add_takeoff(
        "D1", "building", "wall",
        length=10, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )
    p = app.open_project("D1")
    p["boq"][0]["current_quantity"] = 5
    app.store.save("D1", p)
    period = app.save_statement_period("D1", current_quantities={"W001": 5})
    dash = app.financial_dashboard("D1")
    assert dash["project_name"] == "پروژه مالی"
    assert dash["line_count"] == 1
    assert dash["contract_amount"] > 0
    assert dash["cumulative_work"] == period["completed_total"]
    assert dash["statement_count"] == 1
    assert dash["latest_statement_no"] == 1
    assert dash["latest_payable"] == period["payable_current"]
    assert 0 <= dash["progress_percent"] <= 100


def test_project_financial_control_calculates_cost_variance(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("کنترل مالی", "F1")
    app.add_takeoff(
        "F1", "building", "wall",
        length=10, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )
    p = app.open_project("F1")
    p["boq"][0]["current_quantity"] = 5
    app.store.save("F1", p)
    control = app.project_financial_control("F1", planned_cost=5000, actual_cost=4000)
    assert control["earned_value"] == 6000
    assert control["cost_variance"] == 2000
    assert control["schedule_variance"] == pytest.approx(6000 - (5000 * (6000 / 36000)))
    assert control["cost_performance_index"] == pytest.approx(1.5)


def test_project_cost_ledger_persists_and_feeds_financial_control(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("دفتر هزینه", "C1")
    app.add_takeoff(
        "C1", "building", "wall",
        length=10, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )
    app.open_project("C1")["boq"][0]["current_quantity"] = 5
    p = app.open_project("C1")
    p["boq"][0]["current_quantity"] = 5
    app.store.save("C1", p)
    first = app.add_project_cost("C1", "مصالح", 1500, description="خرید سیمان", date="1405/07/09")
    second = app.add_project_cost("C1", "دستمزد", 500)
    summary = app.project_cost_summary("C1")
    assert first["id"] == 1 and second["id"] == 2
    assert summary["entry_count"] == 2
    assert summary["actual_cost"] == 2000
    assert summary["by_category"]["مصالح"] == 1500
    control = app.project_financial_control("C1")
    assert control["actual_cost"] == 2000


def test_project_cost_report_exports_rows_and_category_totals(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("گزارش هزینه", "CR1")
    app.add_project_cost("CR1", "مصالح", 1200, description="بتن", date="1405/07/01")
    app.add_project_cost("CR1", "دستمزد", 800, description="اجرای سقف", date="1405/07/02")
    path = tmp_path / "costs.xlsx"
    result = app.project_cost_report("CR1", "xlsx", path)
    assert result == path
    assert path.exists() and path.stat().st_size > 0
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "مبلغ نهایی" in values
    assert 2000 in values
    assert "مصالح" in values
    assert "دستمزد" in values


def test_financial_dashboard_includes_actual_cost_and_gross_margin(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("داشبورد مالی", "FD1")
    app.add_takeoff(
        "FD1", "building", "wall",
        length=10, width=0.2, height=3,
        price_code="W001", unit_price=1200,
    )
    p = app.open_project("FD1")
    p["boq"][0]["current_quantity"] = 5
    app.store.save("FD1", p)
    app.add_project_cost("FD1", "مصالح", 1500)
    dash = app.financial_dashboard("FD1")
    assert dash["actual_cost"] == 1500
    assert dash["cost_entry_count"] == 1
    assert dash["gross_margin"] == dash["cumulative_work"] - 1500
    assert dash["cost_by_category"]["مصالح"] == 1500


def test_project_receipts_and_financial_position(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("دریافتی", "R1")
    app.add_takeoff("R1", "building", "wall", length=10, width=0.2, height=3, price_code="W001", unit_price=1200)
    app.add_project_cost("R1", "مصالح", 1000)
    app.add_project_receipt("R1", 5000, description="پرداخت کارفرما", reference="REC-1")
    app.add_project_receipt("R1", 2000, description="علی‌الحساب")
    receipts = app.project_receipt_summary("R1")
    assert receipts["entry_count"] == 2
    assert receipts["total_received"] == 7000
    position = app.project_financial_position("R1")
    assert position["received"] == 7000
    assert position["receivable"] == max(position["contract_amount"] - 7000, 0)
    assert position["actual_cost"] == 1000


def test_project_commitment_ledger_and_cash_exposure(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("تعهدات", "C1")
    app.add_project_commitment("C1", 12000, category="پیمانکار", description="قرارداد اجرا", paid_amount=3000)
    app.add_project_commitment("C1", 5000, category="مصالح")
    summary = app.project_commitment_summary("C1")
    assert summary["entry_count"] == 2
    assert summary["committed_total"] == 17000
    assert summary["paid_total"] == 3000
    assert summary["unpaid_total"] == 14000
    app.add_project_receipt("C1", 4000)
    pos = app.project_financial_position("C1")
    assert pos["committed_cost"] == 17000
    assert pos["unpaid_commitments"] == 14000
    assert pos["cash_exposure"] == 13000

def test_project_commitment_rejects_invalid_paid_amount(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("تعهد نامعتبر", "C2")
    with pytest.raises(ValueError):
        app.add_project_commitment("C2", 1000, paid_amount=1200)


def test_project_financial_document_center_links_ledgers_and_summarizes(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("اسناد مالی", "FDOC1")
    commitment = app.add_project_commitment("FDOC1", 12000, category="پیمانکار", paid_amount=2000)
    cost = app.add_project_cost("FDOC1", "پیمانکار", 9000, description="صورت‌حساب پیمانکار")
    receipt = app.add_project_receipt("FDOC1", 5000, description="دریافت کارفرما")
    doc = app.add_project_financial_document(
        "FDOC1", "INV-1405-001", "فاکتور پیمانکار", 9000,
        counterparty="شرکت اجرا", date="1405/07/10", due_date="1405/07/30",
        reference="REF-001", notes="فاکتور مرحله اول", payment_status="partial",
        commitment_id=commitment["id"], cost_entry_id=cost["id"], receipt_id=receipt["id"],
    )
    assert doc["id"] == 1
    assert doc["document_number"] == "INV-1405-001"
    assert doc["commitment_id"] == commitment["id"]
    assert doc["cost_entry_id"] == cost["id"]
    assert doc["receipt_id"] == receipt["id"]
    summary = app.project_financial_document_summary("FDOC1")
    assert summary["document_count"] == 1
    assert summary["document_total"] == 9000
    assert summary["by_payment_status"]["partial"] == 9000
    assert summary["by_type"]["فاکتور پیمانکار"] == 9000
    assert app.project_financial_documents("FDOC1")[0]["counterparty"] == "شرکت اجرا"


def test_project_financial_document_validates_links_and_unique_number(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("اعتبارسنجی سند", "FDOC2")
    with pytest.raises(KeyError):
        app.add_project_financial_document("FDOC2", "INV-1", "فاکتور", 1000, commitment_id=99)
    app.add_project_financial_document("FDOC2", "INV-1", "فاکتور", 1000)
    with pytest.raises(ValueError):
        app.add_project_financial_document("FDOC2", "INV-1", "فاکتور", 500)
    with pytest.raises(ValueError):
        app.add_project_financial_document("FDOC2", "INV-2", "فاکتور", 500, payment_status="unknown")


def test_project_ledger_entries_can_reference_financial_documents(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("ارتباط دفترها", "FDOC3")
    doc = app.add_project_financial_document("FDOC3", "DOC-1", "سند هزینه", 1500)
    cost = app.add_project_cost("FDOC3", "مصالح", 1500, document_id=doc["id"])
    commitment = app.add_project_commitment("FDOC3", 2000, document_id=doc["id"])
    receipt = app.add_project_receipt("FDOC3", 500, document_id=doc["id"])
    assert cost["document_id"] == doc["id"]
    assert commitment["document_id"] == doc["id"]
    assert receipt["document_id"] == doc["id"]


def test_project_financial_report_exports_consolidated_rows(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("گزارش مالی تجمیعی", "FR1")
    commitment = app.add_project_commitment("FR1", 12000, category="پیمانکار", paid_amount=3000, date="1405/07/01", due_date="1405/07/30")
    cost = app.add_project_cost("FR1", "مصالح", 2500, description="بتن", date="1405/07/02")
    receipt = app.add_project_receipt("FR1", 5000, description="پرداخت کارفرما", reference="REC-01")
    app.add_project_financial_document(
        "FR1", "INV-01", "فاکتور", 2500,
        counterparty="پیمانکار", payment_status="partial",
        commitment_id=commitment["id"], cost_entry_id=cost["id"], receipt_id=receipt["id"],
    )
    path = tmp_path / "financial.xlsx"
    result = app.project_financial_report("FR1", "xlsx", path)
    assert result == path
    assert path.exists() and path.stat().st_size > 0
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "نوع" in values
    assert "سند مالی" in values
    assert "تعهد" in values
    assert "هزینه" in values
    assert "دریافتی" in values
    assert 2500 in values
    assert 12000 in values
    assert 5000 in values


def test_project_counterparty_summary_and_report(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("طرف حساب", "CP1")
    app.add_project_commitment("CP1", 12000, category="پیمانکار", paid_amount=3000, counterparty="شرکت الف")
    app.add_project_cost("CP1", "پیمانکار", 2500, description="اجرا", counterparty="شرکت الف")
    app.add_project_receipt("CP1", 5000, description="دریافت", counterparty="کارفرمای اصلی")
    app.add_project_financial_document("CP1", "INV-CP-1", "فاکتور", 2500, counterparty="شرکت الف", payment_status="partial")
    summary = app.project_counterparty_summary("CP1")
    assert summary["counterparty_count"] == 2
    rows = {x["counterparty"]: x for x in summary["counterparties"]}
    assert rows["شرکت الف"]["committed_amount"] == 12000
    assert rows["شرکت الف"]["paid_commitments"] == 3000
    assert rows["شرکت الف"]["unpaid_commitments"] == 9000
    assert rows["شرکت الف"]["actual_cost"] == 2500
    assert rows["شرکت الف"]["document_amount"] == 2500
    assert rows["کارفرمای اصلی"]["received"] == 5000
    path = tmp_path / "counterparties.xlsx"
    assert app.project_counterparty_report("CP1", "xlsx", path) == path
    assert path.exists() and path.stat().st_size > 0


def test_project_counterparty_detailed_ledger(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("گردش حساب", "CP2")
    app.add_project_commitment("CP2", 10000, counterparty="پیمانکار الف", paid_amount=2000, date="1405/07/01")
    app.add_project_cost("CP2", "پیمانکار", 3000, counterparty="پیمانکار الف", date="1405/07/02")
    app.add_project_receipt("CP2", 5000, counterparty="پیمانکار الف", date="1405/07/03")
    app.add_project_financial_document("CP2", "DOC-CP2", "فاکتور", 3000, counterparty="پیمانکار الف", date="1405/07/02")
    ledger = app.project_counterparty_ledger("CP2", "پیمانکار الف")
    assert ledger["row_count"] == 4
    assert ledger["cash_in"] == 5000
    assert ledger["cash_out"] == 13000
    assert ledger["net_cash"] == -8000
    assert [x["source"] for x in ledger["rows"]] == ["تعهد", "سند مالی", "هزینه", "دریافتی"]
    path = tmp_path / "party-ledger.xlsx"
    assert app.project_counterparty_ledger_report("CP2", "پیمانکار الف", "xlsx", path) == path
    assert path.exists() and path.stat().st_size > 0


def test_project_financial_aging_classifies_commitments_and_documents(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("کنترل سررسید", "AGE1")
    app.add_project_commitment(
        "AGE1", 10000, paid_amount=2000, counterparty="پیمانکار الف",
        reference="COM-1", due_date="1405/07/09",
    )
    app.add_project_commitment(
        "AGE1", 5000, paid_amount=5000, counterparty="پیمانکار ب",
        due_date="1405/07/01",
    )
    app.add_project_financial_document(
        "AGE1", "INV-1", "فاکتور", 7000,
        counterparty="فروشنده", due_date="1405/07/09", payment_status="unpaid",
    )
    app.add_project_financial_document(
        "AGE1", "INV-2", "فاکتور", 3000,
        counterparty="فروشنده دوم", due_date="1405/07/15", payment_status="partial",
    )
    app.add_project_financial_document(
        "AGE1", "INV-3", "فاکتور", 2000,
        counterparty="فروشنده سوم", due_date="1405/07/01", payment_status="paid",
    )

    aging = app.project_financial_aging("AGE1", "۱۴۰۵-۰۷-۱۰")
    assert aging["as_of"] == "1405/07/10"
    assert aging["row_count"] == 3
    assert aging["overdue_count"] == 2
    assert aging["due_today_count"] == 0
    assert aging["upcoming_count"] == 1
    assert aging["overdue_commitments"] == 8000
    assert aging["overdue_unpaid_documents"] == 7000
    assert aging["overdue_amount"] == 15000
    assert [x["status"] for x in aging["rows"]] == ["معوق", "معوق", "آتی"]

    path = tmp_path / "aging.xlsx"
    assert app.project_financial_aging_report("AGE1", "1405/07/10", "xlsx", path) == path
    assert path.exists() and path.stat().st_size > 0
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "وضعیت" in values
    assert "معوق" in values
    assert 15000 in values or 8000 in values or 7000 in values


def test_project_financial_aging_rejects_invalid_as_of_date(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("تاریخ نامعتبر", "AGE2")
    with pytest.raises(ValueError):
        app.project_financial_aging("AGE2", "1405/99/99")


def test_project_counterparty_master_registry_normalizes_and_prevents_duplicates(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("دفتر طرف حساب", "CP1")
    first = app.add_project_counterparty("CP1", "  شرکت   نمونه  ", role="پیمانکار", notes="اصلی")
    assert first["id"] == "CP0001"
    assert first["name"] == "شرکت نمونه"
    assert app.find_project_counterparty("CP1", "شرکت نمونه")["id"] == "CP0001"
    with pytest.raises(ValueError):
        app.add_project_counterparty("CP1", "شرکت نمونه")
    second = app.add_project_counterparty("CP1", "فروشنده الف", role="فروشنده")
    assert second["id"] == "CP0002"
    assert [x["name"] for x in app.project_counterparties("CP1")] == ["شرکت نمونه", "فروشنده الف"]


def test_project_financial_alerts_summarize_actionable_aging(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("هشدار مالی", "AL1")
    app.add_project_commitment("AL1", 10000, due_date="۱۴۰۵/۰۷/۰۹")
    app.add_project_commitment("AL1", 5000, due_date="۱۴۰۵/۰۷/۱۰")
    app.add_project_commitment("AL1", 3000)
    alerts = app.project_financial_alerts("AL1", "۱۴۰۵/۰۷/۱۰")
    assert alerts["overdue_amount"] == 10000
    assert alerts["overdue_count"] == 1
    assert alerts["due_today_count"] == 1
    assert alerts["no_due_date_count"] == 1
    assert alerts["has_overdue"] is True
    assert alerts["has_due_today"] is True
    assert alerts["has_missing_due_date"] is True


def test_financial_entries_auto_link_to_counterparty_master(tmp_path):
    app = StructuralProApp(tmp_path)
    app.create_project("اتصال طرف حساب", "LINK1")
    app.add_project_counterparty("LINK1", "شرکت نمونه", role="پیمانکار")
    commitment = app.add_project_commitment("LINK1", 1000, counterparty="  شرکت   نمونه ")
    cost = app.add_project_cost("LINK1", "پیمانکار", 200, counterparty="شرکت نمونه")
    receipt = app.add_project_receipt("LINK1", 300, counterparty="شرکت نمونه")
    document = app.add_project_financial_document("LINK1", "DOC-LINK1", "فاکتور", 500, counterparty="شرکت نمونه")
    assert commitment["counterparty_id"] == "CP0001"
    assert cost["counterparty_id"] == "CP0001"
    assert receipt["counterparty_id"] == "CP0001"
    assert document["counterparty_id"] == "CP0001"
