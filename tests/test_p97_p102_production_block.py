import csv, zipfile
from pathlib import Path
from core.estimate.pricebook_row_extraction_v2 import extract_xlsx
from core.estimate.pricebook_archive_audit_v1 import audit_archive
from core.estimate.pricebook_artifact_matrix_v1 import empty_matrix,validate_matrix,record
from core.estimate.pricebook_scale_integration_v1 import price_rows
from core.iran.full_building_rules_v1 import QuantityRule,validate_registry,DISCIPLINES

def test_p97_multisheet_header_detection(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/"pb.xlsx"; wb=Workbook(); ws=wb.active; ws.title="Cover"
    ws.append(["عنوان فهرست"]); ws.append(["سال 1404"])
    w=wb.create_sheet("فصل 08"); w.append(["شماره ردیف","شرح کار","واحد","بهای واحد"])
    w.append(["080101","بتن","م3",1234]); w.append(["080102","بتن","م3",2345]); wb.save(p)
    rows=extract_xlsx(p,1404,"ابنیه")
    assert {r.item_code for r in rows}=={"080101","080102"}
    assert rows[0].source_sheet=="فصل 08"

def test_p97_cover_only_xlsx_fails_closed(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/"cover_only.xlsx"; wb=Workbook(); ws=wb.active
    ws.title="Cover"; ws.append(["عنوان فهرست"]); ws.append(["سال 1404"]); wb.save(p)
    try:
        extract_xlsx(p,1404,"ابنیه")
    except ValueError as exc:
        assert str(exc)=="no parseable XLSX pricebook rows"
    else:
        raise AssertionError("cover-only workbook must fail closed")

def test_p98_archive_audit_reports_all_files(tmp_path):
    z=tmp_path/"pb.zip"; a=tmp_path/"a.csv"; b=tmp_path/"b.csv"
    a.write_text("ردیف,شرح,واحد,بهای واحد\n100,A,m3,10\n",encoding="utf-8")
    b.write_text("ردیف,شرح,واحد,بهای واحد\n200,B,m,20\n",encoding="utf-8")
    with zipfile.ZipFile(z,"w") as zz: zz.write(a,"a.csv"); zz.write(b,"b.csv")
    rows,manifest=audit_archive(z,1404)
    assert len(rows)==2 and {x.name for x in manifest}=={"a.csv","b.csv"}

def test_p100_matrix():
    m=empty_matrix(); assert not validate_matrix(m)["complete"]
    record(m,1404,"ابنیه",{"sha256":"a"*64})
    assert m[(1404,"ابنیه")]["sha256"]=="a"*64

def test_p101_scaled_pricing():
    class R:
        def __init__(self,c,u,p): self.year=1404; self.discipline="ابنیه"; self.item_code=c; self.unit=u; self.unit_price=p
    out=price_rows([{"takeoff_id":"1","item_code":"۱۲-۳","unit":"m3","quantity":2}],
                   [R("12-3","m3",100)],1404,"ابنیه")
    assert out[0]["total"]==200

def test_p101_unpriced_row_fails_closed():
    class R:
        def __init__(self):
            self.year=1404; self.discipline="ابنیه"; self.item_code="12-3"
            self.unit="m3"; self.unit_price=None
    try:
        price_rows([{"takeoff_id":"1","item_code":"12-3","unit":"m3","quantity":2}],
                   [R()],1404,"ابنیه")
    except ValueError as exc:
        assert str(exc)=="unpriced pricebook row: 12-3/m3"
    else:
        raise AssertionError("unpriced pricebook rows must fail closed")

def test_p102_all_disciplines_contract():
    rules=[QuantityRule(f"r{i}",d,"source","v1",f"c{i}","m","quantity") for i,d in enumerate(DISCIPLINES)]
    assert validate_registry(rules)["complete"]
