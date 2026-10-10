"""Regression tests for fail-closed pricebook replacement and atomic import."""
import csv
import pytest

from core.pricing.catalog import PriceCatalog, PriceItem
from core.pricing.import_service import PricebookImportService


def _item(group, code, price=10):
    return PriceItem(1404, group, "01", code, "test", "m2", price)


def _write(path, rows):
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["year", "group", "chapter", "code", "description", "unit", "unit_price"])
        writer.writeheader()
        for row in rows:
            writer.writerow(dict(year=row.year, group=row.group, chapter=row.chapter, code=row.code,
                                 description=row.description, unit=row.unit, unit_price=row.unit_price))


def test_replace_one_discipline_keeps_other_discipline(tmp_path):
    catalog = PriceCatalog([_item("ابنیه", "100", 10), _item("برق", "200", 20)])
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "300", 30)])
    PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404) is None
    assert catalog.get("200", 1404).unit_price == 20
    assert catalog.get("300", 1404).unit_price == 30


def test_cross_discipline_code_collision_rejected_without_mutation(tmp_path):
    catalog = PriceCatalog([_item("برق", "100", 20)])
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 30)])
    with pytest.raises(ValueError, match="کد یکسان"):
        PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404).group == "برق"
    assert catalog.get("100", 1404).unit_price == 20


def test_duplicate_import_codes_rejected_without_mutation(tmp_path):
    catalog = PriceCatalog([_item("برق", "200", 20)])
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 30), _item("ابنیه", "100", 40)])
    with pytest.raises(ValueError, match="تکراری"):
        PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("200", 1404).unit_price == 20
    assert catalog.get("100", 1404) is None


def test_custom_price_conflict_rejected_without_losing_override(tmp_path):
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    catalog.set_custom_price("100", 1404, 77, reason="approved by user")
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 30)])
    with pytest.raises(ValueError, match="قیمت سفارشی"):
        PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404).unit_price == 77
    assert catalog._items[(1404, "100")].unit_price == 10


def test_unrelated_custom_price_survives_other_discipline_import(tmp_path):
    catalog = PriceCatalog([_item("برق", "200", 20)])
    catalog.set_custom_price("200", 1404, 55)
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 30)])
    PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("200", 1404).unit_price == 55
    assert catalog.get("100", 1404).unit_price == 30


def test_import_preserves_existing_price_history(tmp_path):
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    catalog.add(_item("ابنیه", "100", 15))
    assert catalog.price_history("100", 1404)
    previous = catalog.price_history("100", 1404)
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 30)])
    PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    history = catalog.price_history("100", 1404)
    assert history[:len(previous)] == previous
    assert history[-1] == {
        "year": 1404, "code": "100", "old_price": 15, "new_price": 30,
    }


def test_direct_csv_import_does_not_delete_other_discipline():
    catalog = PriceCatalog([_item("برق", "200", 20), _item("ابنیه", "100", 10)])
    csv_text = PriceCatalog([_item("ابنیه", "300", 30)]).export_csv()
    assert catalog.import_csv(csv_text, replace_year=True) == 1
    assert catalog.get("200", 1404).unit_price == 20
    assert catalog.get("100", 1404) is None
    assert catalog.get("300", 1404).unit_price == 30


def test_direct_csv_import_rejects_conflicting_override_atomically():
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    catalog.set_custom_price("100", 1404, 77)
    csv_text = PriceCatalog([_item("ابنیه", "100", 30)]).export_csv()
    with pytest.raises(ValueError, match="custom prices"):
        catalog.import_csv(csv_text, replace_year=True)
    assert catalog.get("100", 1404).unit_price == 77
    assert catalog._items[(1404, "100")].unit_price == 10


def test_service_import_appends_price_change_to_audit_history(tmp_path):
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 30)])
    PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.price_history("100", 1404)[-1] == {
        "year": 1404, "code": "100", "old_price": 10, "new_price": 30,
    }


def test_replace_rejects_duplicate_legacy_keys_without_mutation():
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    with pytest.raises(ValueError, match="duplicate year/code"):
        catalog.replace([_item("ابنیه", "200", 20), _item("برق", "200", 30)])
    assert catalog.get("100", 1404).unit_price == 10
    assert catalog.get("200", 1404) is None


def test_replace_rejects_invalid_row_atomically():
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    with pytest.raises(ValueError, match="finite"):
        catalog.replace([_item("ابنیه", "200", 20), _item("ابنیه", "300", float("nan"))])
    assert catalog.get("100", 1404).unit_price == 10
    assert catalog.get("200", 1404) is None


def test_replace_protects_custom_prices():
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    catalog.set_custom_price("100", 1404, 77)
    with pytest.raises(ValueError, match="custom prices"):
        catalog.replace([_item("ابنیه", "100", 30)])
    assert catalog.get("100", 1404).unit_price == 77


def test_missing_csv_price_does_not_turn_into_zero_or_modify_catalog(tmp_path):
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    path = tmp_path / "prices.csv"
    path.write_text(
        "year,group,chapter,code,description,unit,unit_price\n"
        "1404,ابنیه,01,100,test,m2,\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="بهای واحد خالی"):
        PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404).unit_price == 10


def test_explicit_zero_csv_price_is_not_treated_as_missing(tmp_path):
    catalog = PriceCatalog()
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "100", 0)])
    PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.resolve("100", 1404)["status"] == "zero_price"


def test_missing_excel_price_is_rejected_without_partial_import(tmp_path):
    from openpyxl import Workbook
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    path = tmp_path / "prices.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["year", "group", "chapter", "code", "description", "unit", "unit_price"])
    ws.append([1404, "ابنیه", "01", "200", "valid", "m2", 20])
    ws.append([1404, "ابنیه", "01", "300", "missing", "m2", None])
    wb.save(path)
    with pytest.raises(ValueError, match="بهای واحد خالی"):
        PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404).unit_price == 10
    assert catalog.get("200", 1404) is None


def test_excel_incomplete_row_is_not_silently_dropped(tmp_path):
    from openpyxl import Workbook
    catalog = PriceCatalog()
    path = tmp_path / "incomplete.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["year", "group", "chapter", "code", "description", "unit", "unit_price"])
    ws.append([1404, "ابنیه", "01", "100", "test", "m2", 10])
    ws.append([1404, "ابنیه", "01", "200", "", "m2", 20])
    wb.save(path)
    with pytest.raises(ValueError, match="ردیف ناقص Excel"):
        PricebookImportService(catalog).import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404) is None


def test_pricebook_modified_after_inspection_fails_without_mutation(tmp_path, monkeypatch):
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    path = tmp_path / "prices.csv"
    _write(path, [_item("ابنیه", "200", 20)])
    service = PricebookImportService(catalog)
    loader = service._load_items
    calls = 0

    def changed_after_loading(*args, **kwargs):
        nonlocal calls
        calls += 1
        result = loader(*args, **kwargs)
        if calls == 2:
            _write(path, [_item("ابنیه", "200", 999)])
        return result

    monkeypatch.setattr(service, "_load_items", changed_after_loading)
    with pytest.raises(ValueError, match="تغییر کرده است"):
        service.import_file(path, year=1404, replace_year=True)
    assert catalog.get("100", 1404).unit_price == 10
    assert catalog.get("200", 1404) is None


def test_direct_catalog_add_rejects_cross_discipline_code_collision():
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    with pytest.raises(ValueError, match="cross-discipline"):
        catalog.add(_item("برق", "100", 20))
    assert catalog.get("100", 1404).group == "ابنیه"
    assert catalog.get("100", 1404).unit_price == 10


def test_direct_catalog_add_does_not_detach_custom_price():
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    catalog.set_custom_price("100", 1404, 77)
    with pytest.raises(ValueError, match="custom prices"):
        catalog.add(_item("ابنیه", "100", 30))
    assert catalog.get("100", 1404).unit_price == 77
    assert catalog._items[(1404, "100")].unit_price == 10


def test_import_rejects_row_year_different_from_selected_year(tmp_path):
    catalog = PriceCatalog([_item("ابنیه", "100", 10)])
    path = tmp_path / "wrong-year.csv"
    _write(path, [PriceItem(1403, "ابنیه", "01", "200", "test", "m2", 20)])
    with pytest.raises(ValueError, match="سال ردیف"):
        PricebookImportService(catalog).import_file(path, year=1404)
    assert catalog.get("100", 1404).unit_price == 10
    assert catalog.get("200", 1403) is None


def test_blank_excel_year_uses_explicit_import_year(tmp_path):
    from openpyxl import Workbook
    catalog = PriceCatalog()
    path = tmp_path / "no-year.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["year", "group", "chapter", "code", "description", "unit", "unit_price"])
    ws.append([None, "ابنیه", "01", "100", "test", "m2", 10])
    wb.save(path)
    PricebookImportService(catalog).import_file(path, year=1404)
    assert catalog.get("100", 1404).unit_price == 10


def test_excel_zero_price_without_row_identity_is_not_skipped(tmp_path):
    from openpyxl import Workbook
    catalog = PriceCatalog()
    path = tmp_path / "zero-without-code.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["year", "group", "chapter", "code", "description", "unit", "unit_price"])
    ws.append([1404, "ابنیه", "01", "100", "test", "m2", 10])
    ws.append([1404, "ابنیه", "01", None, None, None, 0])
    wb.save(path)
    with pytest.raises(ValueError, match="ردیف ناقص Excel"):
        PricebookImportService(catalog).import_file(path, year=1404)
    assert catalog.get("100", 1404) is None
